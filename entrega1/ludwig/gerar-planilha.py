"""Gera entrega1/ludwig/planilha-oraculos-matriz-OpenAPINormalizer.xlsx"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()
bold = Font(bold=True)
hdr = PatternFill("solid", fgColor="D9E1F2")
wrap = Alignment(wrap_text=True, vertical="top")

def sheet(name, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = bold, hdr
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = wrap
    return ws

# ---------------- Resumo ----------------
wb.active.title = "Resumo"
ws = wb.active
ws.append(["Item", "Valor"])
for c in ws[1]:
    c.font, c.fill = bold, hdr
for row in [
    ["Classe sob teste", "org.openapitools.codegen.OpenAPINormalizer (modules/openapi-generator)"],
    ["Baseline", "ef964b04480889ef86b56cfae84ade8ad4c91c41 (v7.25.0)"],
    ["Suíte", "OpenAPINormalizerJUnitTest.java (JUnit 5, @Nested separando por regra, @ParameterizedTest nas famílias)"],
    ["Casos normativos", "46 (CT-NO-01 a CT-NO-46), 47 invocações JUnit (CT-NO-31 roda 2x). Rodada 1: CT-NO-01 a CT-NO-36 (uma regra por vez). Rodada 2, guiada pela cobertura: CT-NO-37 a CT-NO-46 (ramos estruturais)"],
    ["Resultado", "47/47 invocações PASS (46 casos), 0 FAIL — duas rodadas, sem correção em teste ou oráculo"],
    ["Cobertura JaCoCo 0.8.12 (só esta suíte) - classe",
     "45% instruções (2063/4560), 35% desvios (294/847), 46% linhas (512/1109), 62/84 métodos. "
     "O número reflete o recorte em 6 das ~25 regras; os métodos das regras testadas estão quase todos em 100%"],
    ["Métodos sob teste (cobertura de desvios)",
     "getRule, skipNormalization, normalizeArraySchema, normalizeMapSchema e normalizeSimpleSchema 100%; processRemoveXInternalFromOperation 100%; processSetTagsForAllOperations 100%; "
     "processSetTagsToOperationId 100%; processKeepOnlyFirstTagInOperation 7/8; "
     "processSimplifyBooleanEnum 6/8; normalizeOneOf/normalizeAnyOf 7/8; normalizeSchema 35/56"],
    ["Complexidade ciclomática (JaCoCo)", "classe 509; normalizeSchema 29 (é o despachante: array, oneOf, anyOf, allOf, composed, properties, map, boolean, integer, simple)"],
    ["Defeitos", "nenhum observado - nenhuma issue aberta"],
    ["Teste manual", "CTM-NO-01 (spec reemitida pelo gerador openapi via CLI, uma regra por passo + controle): 6/6 PASS"],
]:
    ws.append(row)
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 110
for row in ws.iter_rows():
    for c in row:
        c.alignment = wrap

# ---------------- Oráculos ----------------
oraculos = [
    ["R00", "Processamento de regras",
     "Como as regras são lidas: o construtor recebe (OpenAPI, Map). Por padrão, SIMPLIFY_ONEOF_ANYOF, SIMPLIFY_BOOLEAN_ENUM, "
     "SIMPLIFY_ONEOF_ANYOF_ENUM e REFACTOR_ALLOF_WITH_PROPERTIES_ONLY vêm ligadas e o resto desligado. Cada valor é lido como "
     "string 'true'/'false'. Nome de regra que não existe só gera um WARN no log, sem quebrar nada. DISABLE_ALL faz o normalize() "
     "sair de cara, sem tocar na spec. A fábrica createNormalizer sem NORMALIZER_CLASS devolve OpenAPINormalizer. normalize() cria components/schemas e preenche info (OpenAPI/0.0.1) quando ausentes.",
     "OpenAPINormalizer.java:183-209,240-304,311-326,427-459"],
    ["R01", "SIMPLIFY_BOOLEAN_ENUM",
     "Se um boolean tem enum (com 1 ou 2 valores), o enum é apagado e sobra o boolean puro — um enum booleano não diz nada além do "
     "próprio tipo. Boolean sem enum não muda, e enum de string não é tocado por esta regra.",
     "docs/customization.md (simplifyBooleanEnum); OpenAPINormalizer.java:2131-2142"],
    ["R02", "SIMPLIFY_ONEOF_ANYOF",
     "oneOf/anyOf com um único sub-schema é trocado pelo próprio sub-schema (ex.: SingleAnyOfTest vira string com o enum de 2 itens; "
     "AnyOfTest, que incluía null, vira string nullable). Com a regra desligada (e R03 também), o anyOf fica como está. DISABLE_ALL "
     "mantém tudo intacto.",
     "docs/customization.md (simplifyOneOfAnyOf); OpenAPINormalizer.java:1711-1750"],
    ["R03", "SIMPLIFY_ONEOF_ANYOF_ENUM",
     "Quando todos os sub-schemas de um oneOf/anyOf são enums (unitários ou $refs para enums), eles são fundidos num enum só "
     "(ColorEnum vira [red,green,blue,yellow,purple]; combinedRefsEnum vira [A,B,C,D]). Se o oneOf mistura enum com não-enum, nada é "
     "feito. Com a regra desligada, o oneOf fica como está.",
     "docs/customization.md (simplifyOneOfWithEnums); OpenAPINormalizer.java:1641-1695"],
    ["R04", "KEEP_ONLY_FIRST_TAG_IN_OPERATION",
     "Operação com mais de uma tag fica só com a primeira. Com zero ou uma tag, nada acontece. Com a regra desligada, as tags passam ilesas.",
     "docs/customization.md (enableKeepOnlyFirstTagInOperation); OpenAPINormalizer.java:1468-1479"],
    ["R05", "SET_TAGS_FOR_ALL_OPERATIONS / SET_TAGS_TO_OPERATIONID",
     "SET_TAGS_FOR_ALL_OPERATIONS=<nome> carimba esse nome como tag de todas as operações (string vazia = regra inativa, não mexe em nada). "
     "SET_TAGS_TO_OPERATIONID=true troca as tags pelo operationId de cada operação — e se ela não tem operationId, usa 'default'.",
     "docs/customization.md; OpenAPINormalizer.java:1486-1511"],
    ["R06", "REMOVE_X_INTERNAL",
     "Apaga o x-internal=true das operações, dos schemas em components e das propriedades. Se não há extensions, nada quebra. "
     "Com a regra desligada, o x-internal continua lá.",
     "docs/customization.md (REMOVE_X_INTERNAL); OpenAPINormalizer.java:1448-1460,850-873,1183-1195"],
]
sheet("Oraculos", ["ID", "Regra", "Oráculo", "Fonte"],
      oraculos, [8, 34, 90, 50])

# ---------------- Matriz ----------------
matriz = [
    ["CT-NO-01", "R00", "getRule com Map vazio", "defaults true/false esperados", "PASS"],
    ["CT-NO-02", "R00", "Map com 'false'/'true' explícitos", "regras assumem os valores dados", "PASS"],
    ["CT-NO-03", "R00", "regra inválida + normalize do yaml boolean", "sem exceção; enum apagado", "PASS"],
    ["CT-NO-04", "R00", "DISABLE_ALL sobre yaml boolean", "enum de 2 itens mantido", "PASS"],
    ["CT-NO-05", "R00", "createNormalizer sem NORMALIZER_CLASS", "devolve OpenAPINormalizer", "PASS"],
    ["CT-NO-06", "R01", "BooleanSchema enum [true,false] + normalizeSchema", "enum null", "PASS"],
    ["CT-NO-07", "R01", "yaml BooleanEnumTest + normalize()", "enum null", "PASS"],
    ["CT-NO-08", "R01", "yaml + regra 'false'", "enum de 2 itens mantido", "PASS"],
    ["CT-NO-09", "R01", "BooleanSchema sem enum", "sem exceção; enum null", "PASS"],
    ["CT-NO-10", "R01", "StringSchema enum [a,b]", "enum intacto", "PASS"],
    ["CT-NO-11", "R02", "yaml SingleAnyOfTest (anyOf size 1)", "string, enum size 2, anyOf null", "PASS"],
    ["CT-NO-12", "R02", "yaml AnyOfTest (anyOf com null)", "StringSchema nullable, anyOf null", "PASS"],
    ["CT-NO-13", "R02", "yaml + R02 off + R03 off", "anyOf size 1 mantido", "PASS"],
    ["CT-NO-14", "R02", "yaml + DISABLE_ALL", "anyOf size 1 mantido", "PASS"],
    ["CT-NO-15", "R03", "yaml ColorEnum", "string [red,green,blue,yellow,purple], oneOf null", "PASS"],
    ["CT-NO-16", "R03", "yaml combinedRefsEnum ($refs)", "enum [A,B,C,D], oneOf null", "PASS"],
    ["CT-NO-17", "R03", "yaml + R03 off, parâmetro /test", "oneOf mantido, enum null", "PASS"],
    ["CT-NO-18", "R03", "oneOf misto (string-enum + integer)", "oneOf size 2 mantido", "PASS"],
    ["CT-NO-19", "R04", "op tags [person,basic] + regra", "tags [person]", "PASS"],
    ["CT-NO-20", "R04", "op 1 tag + regra", "tags [person]", "PASS"],
    ["CT-NO-21", "R04", "op sem tags + regra", "sem exceção; tags null", "PASS"],
    ["CT-NO-22", "R04", "op [person,basic] sem regra", "tags intactas", "PASS"],
    ["CT-NO-23", "R05", "op + SET_TAGS_FOR_ALL_OPERATIONS=core", "tags [core]", "PASS"],
    ["CT-NO-24", "R05", "op com/sem operationId + SET_TAGS_TO_OPERATIONID", "[op] / [default]", "PASS"],
    ["CT-NO-25", "R04", "yaml GET [person,basic] + regra", "GET [person]; DELETE size 1", "PASS"],
    ["CT-NO-26", "R06", "op x-internal true + regra", "chave removida", "PASS"],
    ["CT-NO-27", "R06", "op sem extensions + regra", "sem exceção; extensions null", "PASS"],
    ["CT-NO-28", "R06", "op x-internal sem regra", "x-internal mantido", "PASS"],
    ["CT-NO-29", "R06", "schema Dummy x-internal + regra", "chave removida", "PASS"],
    ["CT-NO-30", "R06", "yaml Dummy/DELETE/GET + regra", "removidos; GET null intacto", "PASS"],
    ["CT-NO-31", "R04", "@ParameterizedTest 2 e 3 tags", "só 1ª tag em cada invocação", "PASS (2 invocações)"],
    ["CT-NO-32", "R05", "SET_TAGS_FOR_ALL_OPERATIONS=''", "tags intactas", "PASS"],
    ["CT-NO-33", "R00", "getRule('NO_SUCH_RULE')", "false", "PASS"],
    ["CT-NO-34", "R00", "normalize() sem info", "info OpenAPI/0.0.1 preenchido", "PASS"],
    ["CT-NO-35", "R00", "normalize() sem paths", "sem exceção", "PASS"],
    ["CT-NO-36", "R01", "BooleanSchema enum [true] unitário", "enum null", "PASS"],
]
matriz2 = []
for _c, _r, _e, _x, _res in matriz:
    _obt = "conforme esperado" + (" (2 invocações)" if "2 invoca" in _res else "")
    matriz2.append([_c, _r, _e, _x, _obt, "PASS", "13/10/2026"])
matriz2 += [
    ["CT-NO-37", "R01/R02", "ArraySchema de boolean-enum + normalizeSchema", "array intacto; items sem enum", "conforme esperado", "PASS", "13/10/2026"],
    ["CT-NO-38", "R00", "MapSchema (additionalProperties string)", "map intacto; additionalProperties string", "conforme esperado", "PASS", "13/10/2026"],
    ["CT-NO-39", "R00", "$ref puro sem siblings + normalizeSchema", "$ref intacto (skip)", "conforme esperado", "PASS", "13/10/2026"],
    ["CT-NO-40", "R00", "normalizeSchema(null)", "null, sem exceção", "conforme esperado", "PASS", "13/10/2026"],
    ["CT-NO-41", "R02", "oneOf com os 6 tipos (string, integer, number, boolean, array, object)", "any-type: sem tipo e sem oneOf", "conforme esperado", "PASS", "13/10/2026"],
    ["CT-NO-42", "R06", "propriedade com x-internal em string 'true'", "chave removida", "conforme esperado", "PASS", "13/10/2026"],
    ["CT-NO-43", "R00", "$ref com sibling description + normalizeSchema", "$ref some; vira allOf de 1 item", "conforme esperado", "PASS", "13/10/2026"],
    ["CT-NO-44", "R05", "getRule(SET_TAGS_FOR_ALL_OPERATIONS) com valor dado", "true", "conforme esperado", "PASS", "13/10/2026"],
    ["CT-NO-45", "R00", "normalize() em OpenAPI sem components", "components e schemas criados", "conforme esperado", "PASS", "13/10/2026"],
    ["CT-NO-46", "R04/R05", "op 2 tags + KEEP_FIRST + SET_TAGS=core", "vale o valor fixo: [core]", "conforme esperado", "PASS", "13/10/2026"],
]
sheet("Matriz", ["Caso", "Regra", "Entrada", "Esperado (oráculo prévio)", "Obtido", "Resultado", "Data"],
      matriz2, [10, 8, 44, 44, 30, 10, 12])

# ---------------- Defeitos ----------------
sheet("Defeitos", ["ID", "Casos", "Descrição", "Issue"],
      [["(nenhum)", "—", "46 casos (47 invocações), 47/47 PASS e 0 FAIL. Nenhuma divergência observada; nenhuma issue aberta.",
        "—"]],
      [10, 14, 80, 14])

# ---------------- Fora da suíte ----------------
fora = [
    ["E1", "FILTER / SECURITY_SCHEMES_FILTER", "Têm uma mini-linguagem de filtros (operationId|tag|path|method) com parser próprio. Daria outra entrega sozinhos, então ficaram fora do recorte."],
    ["E2", "NORMALIZE_31SPEC", "É a normalização do OpenAPI 3.1 (tipos, exclusiveMin/Max numéricos). Exigiria fixtures 3.1 dedicadas, fora do recorte 3.0 dos meus casos."],
    ["E3", "REPLACE_ONE_OF_BY_DISCRIMINATOR_MAPPING", "Mexe com discriminator/mapping e herança entre schemas; o oráculo do 'certo' acaba dependendo do gerador escolhido."],
    ["E4", "ADD_UNSIGNED_TO_INTEGER_WITH_INVALID_MAX_VALUE", "Já coberta pelo teste upstream; os valores-limite 2^32-1/2^64-1 mereceriam casos próprios; ficou fora do recorte R00–R06 por tempo."],
    ["E5", "REFACTOR_ALLOF_WITH_PROPERTIES_ONLY / resto do allOf", "Vem ligado por padrão, mas não montei oráculo dedicado: composição allOf com/sem properties no mesmo nível fica como trabalho futuro."],
    ["E6", "SORT_MODEL_PROPERTIES", "Ordenação de propriedades; ortogonal ao recorte, sem casos meus."],
    ["E7", "SET_CONTAINER_TO_NULLABLE / SET_PRIMITIVE_TYPES_TO_NULLABLE", "Regras de nullable por categoria; ortogonais ao recorte, sem casos meus."],
    ["E8", "FIX_DUPLICATED_OPERATIONID", "Renomeia operationIds duplicados; ortogonal ao recorte, sem casos meus."],
]
sheet("Fora-da-suite", ["ID", "Tópico", "Motivo"], fora, [8, 48, 80])

wb.save("entrega1/ludwig/planilha-oraculos-matriz-OpenAPINormalizer.xlsx")
print("OK - planilha gerada")
